"""Extract the canonical palika/ward reference data from JE-083_84.xlsx.

This is the single source of truth for everything downstream:
  * KoBo form choice lists (palika, and wards cascaded from palika)
  * BigQuery ref_local_level / ref_ward seed rows
  * The dashboard's coverage denominators

The workbook is the approved district microplan, so it is read-only input and is
never modified. Run:  python scripts/extract_reference.py
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "JE-083_84.xlsx"
OUT_JSON = ROOT / "data" / "reference.json"

TOTAL_SHEET = "Total "          # trailing space is in the real file
JE_TARGET_RATIO = 0.48          # 30+ share of population used across the microplan


def slug(name: str) -> str:
    """KoBo choice names allow only letters, numbers and underscore."""
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", ascii_name.lower())).strip("_")


def header_offset(ws) -> int:
    """Most sheets carry a title row before the 'ward no' header; Birtamod does not.

    Returns the row number of the 'ward no' header so ward rows can be located
    without assuming a fixed layout.
    """
    for row in ws.iter_rows(min_row=1, max_row=6):
        first = row[0].value
        if isinstance(first, str) and first.strip().lower().startswith("ward no"):
            return row[0].row
    raise ValueError(f"no 'ward no' header found in sheet {ws.title!r}")


def read_local_level(ws) -> dict:
    """Read one local-level sheet into {wards: [...], sheet_total: (pop, target)}."""
    start = header_offset(ws) + 3        # header, age-band row, sex row, then data
    wards, sheet_total = [], None

    for row in ws.iter_rows(min_row=start):
        cells = [c.value for c in row]
        key, pop, target = cells[0], cells[1], cells[2]

        if isinstance(key, (int, float)) and not isinstance(key, bool):
            if pop is None:
                raise ValueError(f"{ws.title} ward {key}: missing population")
            wards.append({"ward_no": int(key), "population": int(pop),
                          "je_target": round(float(target), 2) if target is not None else None})
        elif isinstance(key, str) and key.strip().lower() == "total":
            # Buddhashanti and Arjundhara have a blank Total row - tolerated,
            # because the ward rows are complete and the Total sheet has the figure.
            if pop is not None:
                sheet_total = (int(pop), round(float(target), 2))

    return {"wards": wards, "sheet_total": sheet_total}


def read_district_totals(ws) -> dict:
    """The Total sheet is authoritative for each local level's population."""
    totals = {}
    for row in ws.iter_rows(min_row=5):
        name, pop, target = (row[0].value, row[1].value, row[2].value)
        if not isinstance(name, str) or pop is None:
            continue
        if name.strip().lower() == "total":
            continue
        totals[name.strip()] = {"population": int(pop), "je_target": round(float(target), 2)}
    return totals


def main() -> int:
    if not WORKBOOK.exists():
        sys.exit(f"workbook not found: {WORKBOOK}")

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    district_totals = read_district_totals(wb[TOTAL_SHEET])

    local_levels, problems = [], []

    for sheet_name in wb.sheetnames:
        if sheet_name == TOTAL_SHEET:
            continue

        name = sheet_name.strip()
        parsed = read_local_level(wb[sheet_name])
        wards = parsed["wards"]

        ward_pop = sum(w["population"] for w in wards)
        ward_target = round(sum(w["je_target"] for w in wards), 2)

        # The Total sheet is authoritative; ward rows must reconcile to it.
        authoritative = district_totals.get(name)
        if authoritative is None:
            problems.append(f"{name}: absent from the Total sheet")
        elif authoritative["population"] != ward_pop:
            problems.append(
                f"{name}: wards sum to {ward_pop:,} but Total sheet says "
                f"{authoritative['population']:,}"
            )
        elif abs(authoritative["je_target"] - ward_target) > 0.5:
            problems.append(
                f"{name}: ward targets sum to {ward_target:,.2f} but Total sheet says "
                f"{authoritative['je_target']:,.2f}"
            )

        local_levels.append({
            "code": slug(name),
            "name": name,
            "population": ward_pop,
            "je_target": ward_target,
            "ward_count": len(wards),
            "wards": [
                {
                    "code": f"{slug(name)}_{w['ward_no']}",
                    "ward_no": w["ward_no"],
                    "population": w["population"],
                    "je_target": w["je_target"] if w["je_target"] is not None
                    else round(w["population"] * JE_TARGET_RATIO, 2),
                }
                for w in wards
            ],
        })

    if problems:
        print("RECONCILIATION FAILED - refusing to write reference data:", file=sys.stderr)
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        return 1

    district = {
        "district": "Jhapa",
        "fiscal_year": "2083/84",
        "campaign_start": "2026-10-05",
        "campaign_end": "2026-10-19",
        "population": sum(ll["population"] for ll in local_levels),
        "je_target": round(sum(ll["je_target"] for ll in local_levels), 2),
        "local_level_count": len(local_levels),
        "ward_count": sum(ll["ward_count"] for ll in local_levels),
        "local_levels": local_levels,
    }

    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(district, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"wrote {OUT_JSON.relative_to(ROOT)}")
    print(f"  local levels : {district['local_level_count']}")
    print(f"  wards        : {district['ward_count']}")
    print(f"  population   : {district['population']:,}")
    print(f"  JE target    : {district['je_target']:,.2f}  (displays as {round(district['je_target']):,})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
